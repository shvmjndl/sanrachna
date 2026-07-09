"""Multi-format document converters."""

from __future__ import annotations

import email
import re
from email import policy
from pathlib import Path


def convert_to_markdown(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    converters = {
        ".txt": _read_text,
        ".md": _read_text,
        ".csv": _read_text,
        ".eml": _convert_eml,
        ".docx": _convert_docx,
        ".pptx": _convert_pptx,
    }
    fn = converters.get(ext)
    if not fn:
        return f"[Unsupported format: {ext}]"
    return fn(file_path)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _convert_eml(path: Path) -> str:
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)

    parts = [
        f"# Email: {msg.get('subject', 'No Subject')}",
        f"**From:** {msg.get('from', 'Unknown')}",
        f"**To:** {msg.get('to', 'Unknown')}",
        f"**Date:** {msg.get('date', 'Unknown')}",
        "",
    ]

    body = msg.get_body(preferencelist=("plain", "html"))
    if body:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            content = re.sub(r"<[^>]+>", " ", content)
        parts.append(content.strip())

    for part in msg.iter_attachments():
        parts.append(f"\n**Attachment:** {part.get_filename() or 'unnamed'}")

    return "\n".join(parts)


def _convert_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else ""
        if "Heading" in style:
            level = style.replace("Heading ", "")
            parts.append(f"{'#' * int(level) if level.isdigit() else '##'} {text}")
        else:
            parts.append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
        if rows:
            header_sep = "| " + " | ".join("---" for _ in table.rows[0].cells) + " |"
            parts.append("\n".join([rows[0], header_sep] + rows[1:]))

    return "\n\n".join(parts)


def _convert_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"## Slide {i}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    rows.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
                if rows:
                    sep = "| " + " | ".join("---" for _ in table.rows[0].cells) + " |"
                    parts.append("\n".join([rows[0], sep] + rows[1:]))

    return "\n\n".join(parts)
